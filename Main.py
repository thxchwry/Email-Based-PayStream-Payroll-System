import openpyxl
import os
import ezgmail
import pikepdf
from datetime import datetime
# ดึงฟังก์ชันมาจากอีกไฟล์
from create_pay_slip import generate_pdf

# ==========================================
# ตั้งค่าพื้นฐาน
# ==========================================
EXCEL_FILE = 'slip_report_updated.xlsx'
DOMAIN = "@rsu.ac.th"

# รายชื่อเดือน
thai_months = ["มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน", 
               "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]
en_months = ["January", "February", "March", "April", "May", "June", 
             "July", "August", "September", "October", "November", "December"]

now = datetime.now()
current_month_thai = thai_months[now.month - 1]
current_year_be = now.year + 543
time_label = f"{current_month_thai} {current_year_be}" 

print("--- RSU Slip Delivery System Starting ---")
print("Initialing Gmail...")
ezgmail.init() 

# โหลด Excel
wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
sheet = wb.active

# เริ่มวนลูปอ่านพนักงานตั้งแต่แถวที่ 4
for row in range(4, sheet.max_row + 1):
    emp_id = sheet.cell(row=row, column=1).value
    first_name = str(sheet.cell(row=row, column=2).value).strip() 
    last_name = str(sheet.cell(row=row, column=3).value).strip()  
    start_date = sheet.cell(row=row, column=4).value              
    
    salary_val = sheet.cell(row=row, column=5).value or 0
    net_val = sheet.cell(row=row, column=13).value or 0

    if not first_name or first_name == 'None':
        continue

    # ---จัดการวันที่และรหัสผ่าน ---
    if isinstance(start_date, datetime):
        pass_str = start_date.strftime('%Y%m%d')
        start_year_be = start_date.year 
    else:
        date_str = str(start_date)
        pass_str = date_str.replace('-', '')
        start_year_be = date_str.split('-')[0]

    short_year = str(start_year_be)[-2:]
    email_target = f"{first_name.lower()}.{last_name[0].lower()}{short_year}{DOMAIN}"
    
    # ---จัดการโครงสร้าง Folder (ปี > ชื่อพนักงาน) ---
    # สร้างเส้นทาง: ปี_2569/Thidapond
    base_folder = f"ปี_{current_year_be}"
    user_folder = os.path.join(base_folder, first_name)
    
    # สร้าง Folder แบบซ้อนกัน (ถ้าไม่มีให้สร้าง ถ้ามีแล้วไม่ทำอะไร)
    os.makedirs(user_folder, exist_ok=True)

    # ---สร้างไฟล์ PDF และล็อกรหัสผ่าน ---
    temp_pdf = "temp_unlocked.pdf"
    file_name = f"สลิป_{current_month_thai}_{current_year_be}.pdf"
    
    # แก้จุดนี้: ให้ Path ชี้ไปที่ user_folder (Folder ชื่อพนักงาน)
    final_pdf_path = os.path.join(user_folder, file_name)
    
    # สร้างเนื้อหา PDF
    time_label_en = f"{en_months[now.month - 1]} {current_year_be}"
    generate_pdf(temp_pdf, f"{first_name} {last_name}",time_label_en ,salary_val, net_val)

    # ล็อกรหัสผ่าน PDF และบันทึกลงใน Folder พนักงานโดยตรง
    with pikepdf.open(temp_pdf) as p_pdf:
        p_pdf.save(final_pdf_path, encryption=pikepdf.Encryption(owner=pass_str, user=pass_str))
    
    # ลบไฟล์ชั่วคราวทิ้งทันที
    if os.path.exists(temp_pdf):
        os.remove(temp_pdf)

    # ---ส่งอีเมล ---
    try:
        subject = f"🌟 OFFICIAL REMITTANCE ADVICE | ผลตอบแทนแห่งความสำเร็จประจำเดือน {current_month_thai} - {first_name.upper()}"

        body = (f"เรียน คุณ {first_name} ผู้เป็นพลังขับเคลื่อนอันล้ำค่าของ LAE THAE 99,\n\n"
        f"เพราะความทุ่มเทและศักยภาพอันโดดเด่นของท่านคือหัวใจสำคัญที่ทำให้ LAE THAE 99 เติบโตอย่างสง่างามในทุกย่างก้าว "
        f"เราจึงมีความภาคภูมิใจอย่างยิ่งที่จะส่งมอบ 'รายงานสรุปผลตอบแทนแห่งความสำเร็จ (Exclusive Salary Advice)' "
        f"ประจำงวดเดือน {current_month_thai} {current_year_be} เพื่อเป็นการขอบคุณในทุกความมุ่งมั่นที่ท่านมอบให้\n\n"
        f"ขอให้ทุกตัวเลขในเอกสารฉบับนี้ เป็นแรงบันดาลใจให้ท่านก้าวสู่ความสำเร็จที่ยิ่งใหญ่กว่าเดิมร่วมกับเรา\n\n"
        f"--------------------------------------------------\n"
        f"🔐 เอกสารแนบนี้ได้รับการเข้ารหัสลับเพื่อความเป็นส่วนตัวขั้นสูงสุดของท่าน\n"
        f"--------------------------------------------------\n\n"
        f"ขอบคุณที่เป็นหนึ่งในฟันเฟืองที่สมบูรณ์แบบที่สุดของเรา\n\n"
        f"ด้วยความเคารพอย่างสูง,\n"
        f"LAE THAE 99 CO., LTD.")

        
        print(f"[{emp_id}] ส่งหา: {email_target} (เก็บที่: {final_pdf_path})")
        # ส่งไฟล์ที่อยู่ใน Folder พนักงาน
        ezgmail.send(email_target, subject, body, [final_pdf_path])
        print("-> ส่งอีเมลสำเร็จ!")
    except Exception as e:
        print(f"-> พลาด: {e}")

print("\n--- เสร็จสมบูรณ์! ไฟล์ทั้งหมดถูกเก็บเข้า Folder ชื่อพนักงานเรียบร้อยแล้ว ---")